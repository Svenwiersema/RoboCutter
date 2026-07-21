from docx import Document

class RoboDocument:
    VERSION="1.0"

    def __init__(self,title="",code="",version="1.0"):
        self.doc=Document()
        self.title=title
        self.code=code
        self.version=version

    def add_cover(self):
        self.doc.add_heading(self.title,0)
        self.doc.add_paragraph(f"Code: {self.code}")
        self.doc.add_paragraph(f"Versie: {self.version}")

    def add_heading(self,text,level=1):
        self.doc.add_heading(text,level)

    def add_paragraph(self,text):
        self.doc.add_paragraph(text)

    def save(self,path):
        self.doc.save(path)
